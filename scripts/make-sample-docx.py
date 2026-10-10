# 生成示例 .docx 文档,用于测试 Word 导入(每行 "单词 - 释义")
import os
from docx import Document

doc = Document()
doc.add_heading('英语词汇表(Word 示例)', level=1)

items = [
    ("apple", "苹果"),
    ("banana", "香蕉"),
    ("perseverance", "毅力;坚持不懈"),
    ("gratitude", "感激;感恩"),
    ("opportunity", "机会;机遇"),
    ("challenge", "挑战"),
    ("achieve", "实现;达成"),
    ("improve", "改善;提高"),
    ("remember", "记住;记得"),
    ("delicious", "美味的"),
    ("beautiful", "美丽的"),
    ("important", "重要的"),
    ("Good morning, everyone.", "大家早上好。"),
    ("Practice makes perfect.", "熟能生巧。"),
    ("Never give up on your dreams.", "永远不要放弃你的梦想。"),
]

for term, meaning in items:
    doc.add_paragraph(f"{term} - {meaning}")

out = os.path.join(os.path.dirname(__file__), "..", "web", "sample", "示例单词表.docx")
doc.save(out)
print("saved:", os.path.abspath(out))
